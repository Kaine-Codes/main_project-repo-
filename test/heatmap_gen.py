import pandas as pd
import numpy as np
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

import matplotlib.pyplot as plt


# ============================================================
# SETTINGS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

INPUT_FILE = (
    BASE_DIR
    / "output"
    / "chb08_05_features_4s.xlsx"
    
)

OUTPUT_FILE = (
    BASE_DIR
    / "output"
    / "EEG_Channel_Stage_Analysis_05_4s.xlsx"
)

SHEET_NAME = "Features_Long"


# ============================================================
# YOUR 8 FEATURES
# ============================================================

FEATURES = [
    "Hjorth_Activity",
    "Hjorth_Mobility",
    "Hjorth_Complexity",
    "ZCR",
    "Alpha_Theta",
    "Higuchi_FD",
    "Gamma_Beta",
    "Theta_Beta"
]


STAGES = [
    "interictal",
    "preictal",
    "ictal",
    "postictal"
]


# ============================================================
# LOAD DATA
# ============================================================

print("=" * 60)
print(" EEG CHANNEL × SEIZURE STAGE ANALYSIS")
print("=" * 60)

print("\nLoading Features_Long...")

df = pd.read_excel(
    INPUT_FILE,
    sheet_name=SHEET_NAME
)

print(f"Rows loaded: {len(df):,}")

print("\nColumns:")
print(df.columns.tolist())


# ============================================================
# CLEAN STAGE NAMES
# ============================================================

df["segment"] = (
    df["segment"]
    .astype(str)
    .str.strip()
    .str.lower()
)


# Keep only the four stages
df = df[
    df["segment"].isin(STAGES)
].copy()


# ============================================================
# CLEAN CHANNEL NAMES
# ============================================================

df["channel"] = (
    df["channel"]
    .astype(str)
    .str.strip()
)


# ============================================================
# CONVERT FEATURE COLUMNS TO NUMERIC
# ============================================================

for feature in FEATURES:

    df[feature] = pd.to_numeric(
        df[feature],
        errors="coerce"
    )


# ============================================================
# SHOW BASIC INFORMATION
# ============================================================

print("\nStages found:")

print(
    df["segment"]
    .value_counts()
)


print("\nChannels found:")

print(
    df["channel"]
    .unique()
)


print(
    f"\nNumber of channels: "
    f"{df['channel'].nunique()}"
)


# ============================================================
# OUTPUT DIRECTORY
# ============================================================

PNG_DIR = (
    BASE_DIR
    / "output"
    / "EEG_Heatmaps"
)

PNG_DIR.mkdir(
    exist_ok=True
)


# ============================================================
# STORE RESULTS
# ============================================================

median_tables = {}
change_tables = {}


# ============================================================
# PROCESS EACH FEATURE
# ============================================================

for feature in FEATURES:

    print("\n" + "-" * 60)
    print(f"Processing: {feature}")
    print("-" * 60)

    # --------------------------------------------------------
    # MEDIAN VALUE
    #
    # channel × stage
    # --------------------------------------------------------

    table = (
        df
        .groupby(
            ["channel", "segment"]
        )[feature]
        .median()
        .unstack()
    )

    # Force stage order
    table = table.reindex(
        columns=STAGES
    )

    median_tables[feature] = table


    # ========================================================
    # PRINT TABLE
    # ========================================================

    print("\nMedian values:")

    print(
        table.round(4)
    )


    # ========================================================
    # CREATE NORMALIZED CHANGE TABLE
    #
    # Relative to INTERICTAL
    # ========================================================

    baseline = table["interictal"]

    change = pd.DataFrame(
        index=table.index
    )

    for stage in [
        "preictal",
        "ictal",
        "postictal"
    ]:

        # Percentage change
        change[stage] = (
            (
                table[stage] - baseline
            )
            /
            baseline.replace(
                0,
                np.nan
            )
        ) * 100

    change_tables[feature] = change


    # ========================================================
    # CREATE MEDIAN HEATMAP
    # ========================================================

    plt.figure(
        figsize=(9, 8)
    )

    plt.imshow(
        table.values,
        aspect="auto",
        interpolation="nearest"
    )

    plt.colorbar(
        label=feature
    )

    plt.xticks(
        range(len(table.columns)),
        [
            x.capitalize()
            for x in table.columns
        ]
    )

    plt.yticks(
        range(len(table.index)),
        table.index
    )

    plt.xlabel(
        "Seizure Stage"
    )

    plt.ylabel(
        "EEG Channel"
    )

    plt.title(
        f"{feature}\n"
        "Median value by channel and seizure stage"
    )

    plt.tight_layout()

    plt.savefig(
        PNG_DIR
        / f"{feature}_stage_heatmap.png",
        dpi=200
    )

    plt.close()


    # ========================================================
    # CREATE CHANGE HEATMAP
    # ========================================================

    plt.figure(
        figsize=(8, 8)
    )

    max_abs = np.nanmax(
        np.abs(change.values)
    )

    plt.imshow(
        change.values,
        aspect="auto",
        interpolation="nearest",
        cmap="RdBu_r",
        vmin=-max_abs,
        vmax=max_abs
    )

    plt.colorbar(
        label="% change from interictal"
    )

    plt.xticks(
        range(len(change.columns)),
        [
            x.capitalize()
            for x in change.columns
        ]
    )

    plt.yticks(
        range(len(change.index)),
        change.index
    )

    plt.xlabel(
        "Seizure Stage"
    )

    plt.ylabel(
        "EEG Channel"
    )

    plt.title(
        f"{feature}\n"
        "Change relative to interictal"
    )

    plt.axvline(
        -0.5,
        linewidth=1
    )

    plt.tight_layout()

    plt.savefig(
        PNG_DIR
        / f"{feature}_change_heatmap.png",
        dpi=200
    )

    plt.close()


# ============================================================
# CREATE COMBINED PREICTAL TABLE
# ============================================================

print("\n" + "=" * 60)
print("Creating combined preictal analysis...")
print("=" * 60)


preictal_combined = pd.DataFrame()


for feature in FEATURES:

    preictal_combined[
        feature
    ] = change_tables[
        feature
    ]["preictal"]


# ============================================================
# ABSOLUTE PREICTAL CHANGE
# ============================================================

preictal_abs = (
    preictal_combined
    .abs()
)


# ============================================================
# AVERAGE ABSOLUTE CHANGE ACROSS FEATURES
# ============================================================

preictal_summary = pd.DataFrame(
    index=preictal_combined.index
)

preictal_summary[
    "Mean_Absolute_Preictal_Change_%"
] = (
    preictal_abs.mean(
        axis=1
    )
)

preictal_summary[
    "Median_Absolute_Preictal_Change_%"
] = (
    preictal_abs.median(
        axis=1
    )
)

preictal_summary[
    "Features_with_>20%_Change"
] = (
    preictal_abs > 20
).sum(axis=1)


preictal_summary = (
    preictal_summary
    .sort_values(
        "Mean_Absolute_Preictal_Change_%",
        ascending=False
    )
)


# ============================================================
# CREATE EXCEL FILE
# ============================================================

print("\nWriting Excel file...")

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    # --------------------------------------------------------
    # Median tables
    # --------------------------------------------------------

    for feature in FEATURES:

        sheet_name = (
            f"Median_{feature}"
        )[:31]

        median_tables[
            feature
        ].to_excel(
            writer,
            sheet_name=sheet_name
        )


    # --------------------------------------------------------
    # Preictal change
    # --------------------------------------------------------

    preictal_combined.to_excel(
        writer,
        sheet_name="Preictal_Change"
    )


    # --------------------------------------------------------
    # Preictal summary
    # --------------------------------------------------------

    preictal_summary.to_excel(
        writer,
        sheet_name="Preictal_Summary"
    )


    # --------------------------------------------------------
    # Ictal change
    # --------------------------------------------------------

    ictal_combined = pd.DataFrame()

    for feature in FEATURES:

        ictal_combined[
            feature
        ] = change_tables[
            feature
        ]["ictal"]

    ictal_combined.to_excel(
        writer,
        sheet_name="Ictal_Change"
    )


    # --------------------------------------------------------
    # Postictal change
    # --------------------------------------------------------

    postictal_combined = pd.DataFrame()

    for feature in FEATURES:

        postictal_combined[
            feature
        ] = change_tables[
            feature
        ]["postictal"]

    postictal_combined.to_excel(
        writer,
        sheet_name="Postictal_Change"
    )


# ============================================================
# FORMAT EXCEL
# ============================================================

print("Formatting Excel...")

wb = load_workbook(
    OUTPUT_FILE
)


for ws in wb.worksheets:

    # --------------------------------------------------------
    # Freeze headers
    # --------------------------------------------------------

    ws.freeze_panes = "B2"


    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )


    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    ws.column_dimensions[
        "A"
    ].width = 20

    for col in range(
        2,
        ws.max_column + 1
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = 20


    # --------------------------------------------------------
    # Heatmap formatting
    # --------------------------------------------------------

    if (
        ws.max_row > 1
        and ws.max_column > 1
    ):

        start = "B2"

        end = (
            f"{get_column_letter(ws.max_column)}"
            f"{ws.max_row}"
        )

        cell_range = (
            f"{start}:{end}"
        )


        # Diverging heatmap
        # useful for change tables

        if (
            "Change" in ws.title
            or "Summary" in ws.title
        ):

            rule = ColorScaleRule(

                start_type="min",
                start_color="F4CCCC",

                mid_type="num",
                mid_value=0,
                mid_color="FFFFFF",

                end_type="max",
                end_color="C9DAF8"
            )

        else:

            rule = ColorScaleRule(

                start_type="min",
                start_color="FFFFFF",

                mid_type="percentile",
                mid_value=50,
                mid_color="FFF2CC",

                end_type="max",
                end_color="F4CCCC"
            )


        ws.conditional_formatting.add(
            cell_range,
            rule
        )


wb.save(
    OUTPUT_FILE
)


# ============================================================
# FINAL OUTPUT
# ============================================================

print("\n" + "=" * 60)
print("DONE!")
print("=" * 60)

print(
    f"\nExcel file:\n{OUTPUT_FILE}"
)

print(
    f"\nPNG heatmaps:\n{PNG_DIR}"
)

print("\nGenerated:")
print("  ✓ 8 median channel × stage heatmaps")
print("  ✓ 8 interictal-relative change heatmaps")
print("  ✓ Preictal_Change")
print("  ✓ Preictal_Summary")
print("  ✓ Ictal_Change")
print("  ✓ Postictal_Change")